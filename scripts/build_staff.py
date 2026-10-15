"""Normalize the PEN staff workbook into private/staff.json.

Each sheet is its own population (different schools and districts), so unlike
the student census none of them is dropped as a duplicate extract.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT.parent / "PEN" / "Staff Data PEIMA, PSRP-I,II, ESP, Adopted Punjab, Sindh.xlsx"
OUT = ROOT / "private" / "staff.json"

PLACE_ALIASES = {
    "pak pattan": "Pakpattan",
    "tmk": "Tando Muhammad Khan",
    "tmk ii": "Tando Muhammad Khan",
    "dariya khan": "Darya Khan",
    "darya khan": "Darya Khan",
    "mankara": "Mankera",
    "mankera": "Mankera",
    "esakhel": "Isa Khel",
    "isa khel": "Isa Khel",
    "shuja abad": "Shujabad",
    "bulri shah karim": "Bulri Shah Karim",
    "jarwanwala": "Jarwanwala",
}

RELIGION = {
    "islam": "Islam",
    "lslam": "Islam",
    "isalm": "Islam",
    "muslim": "Islam",
    "christian": "Christian",
    "christianity": "Christian",
    "hindu": "Hindu",
    "hindusim": "Hindu",
    "hinduism": "Hindu",
}

SHEETS = [
    {
        "name": "PEIMA Staff data",
        "source": "PEIMA staff",
        "program": "PEIMA",
        "system": "PEIMA",
        "note": "PEIMA schools. Counted on its own; it is not a copy of another staff sheet.",
    },
    {
        "name": "PSRP-I Staff Data",
        "source": "PSRP-I staff",
        "program": "PSRP-I",
        "system": "PEF",
        "note": "PSRP-I partner schools.",
    },
    {
        "name": " PSRP-II Staff Data",
        "source": "PSRP-II staff",
        "program": "PSRP-II",
        "system": "PEF",
        "note": "PSRP-II partner schools. The sheet name has a leading space in the workbook.",
    },
    {
        "name": "Sindh Staff Data",
        "source": "Sindh PSP staff",
        "program": "PSP",
        "system": "PSP",
        "note": "Sindh Education Foundation PSP schools. This sheet has no religion column. District code TMK is Tando Muhammad Khan.",
    },
    {
        "name": "Punjab Staff Data",
        "source": "Punjab adopted staff",
        "program": "PEN Adoption",
        "system": "PEN",
        "note": "Adopted Punjab schools. This sheet has no year column.",
    },
    {
        "name": "ESP Staff Data",
        "source": "Evening School Program staff",
        "program": "Evening School Program",
        "system": "PEN",
        "note": "PEN evening schools. A blank project cell still belongs to this sheet.",
    },
]

PROGRAMS = {
    "peima": ("PEIMA", "PEIMA"),
    "psrp-1": ("PSRP-I", "PEF"),
    "psrp-i": ("PSRP-I", "PEF"),
    "psrp 1": ("PSRP-I", "PEF"),
    "psrp-2": ("PSRP-II", "PEF"),
    "psrp-ii": ("PSRP-II", "PEF"),
    "psrp 2": ("PSRP-II", "PEF"),
    "sindh": ("PSP", "PSP"),
    "psp": ("PSP", "PSP"),
    "sef": ("PSP", "PSP"),
    "esp": ("Evening School Program", "PEN"),
    "evening": ("Evening School Program", "PEN"),
    "evening school program": ("Evening School Program", "PEN"),
    "pen": ("PEN Adoption", "PEN"),
    "pen adoption": ("PEN Adoption", "PEN"),
}

DONORS = {
    "govt": "Govt",
    "govt.": "Govt",
    "government": "Govt",
    "sef": "SEF",
    "pen": "PEN",
    "peima": "PEIMA",
}

ROLE_RULES = [
    ("vice principal", "Vice Principal", "Leadership"),
    ("head teacher", "Head Teacher", "Leadership"),
    ("second head", "Second Head", "Leadership"),
    ("sec head", "Second Head", "Leadership"),
    ("principal", "Principal", "Leadership"),
    ("assistant teacher", "Assistant Teacher", "Teaching"),
    ("library teacher", "Library Teacher", "Teaching"),
    ("sports teacher", "Sports Teacher", "Teaching"),
    ("teacher", "Teacher", "Teaching"),
    ("admin coordinator", "Admin Coordinator", "Administration"),
    ("admin officer", "Admin Officer", "Administration"),
    ("admin assistant", "Admin Assistant", "Administration"),
    ("cluster manager", "Cluster Manager", "Administration"),
    ("m&e", "M&E Officer", "Administration"),
    ("me officer", "M&E Officer", "Administration"),
    ("admission officer", "Admission Officer", "Administration"),
    ("coordinat", "Coordinator", "Administration"),
    ("sweeper", "Sweeper", "Support"),
    ("guard", "Guard", "Support"),
    ("gardener", "Gardener", "Support"),
    ("gardner", "Gardener", "Support"),
    ("office boy", "Office Boy", "Support"),
    ("driver", "Driver", "Support"),
    ("peon", "Peon", "Support"),
    ("aya", "Aya", "Support"),
]

MASTERS = {"masters", "master", "msc", "mcs", "mphil", "mphill", "ma", "ms", "med", "mba"}
BACHELORS = {"bachelors", "bachelor", "ba", "bsc", "bcom", "bs", "bed", "ad", "ads", "ada", "ade"}
INTER = {"inter", "intermediate", "fa", "fsc", "ics", "icom"}
NONE = {"na", "none", "nil", "nill"}
ILLITERATE = {"illiterate", "illitrate", "illiteracy"}


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def norm_header(value) -> str:
    return re.sub(r"\s+", " ", text(value).lower())


def digits(value) -> str:
    return re.sub(r"\D", "", text(value))


def pretty(value: str) -> str:
    token = text(value)
    if not token:
        return "Unknown"
    letters = re.sub(r"[^A-Za-z]", "", token)
    if not letters or not (letters.isupper() or letters.islower()):
        return token
    small = {"of", "and", "the", "al"}
    parts = []
    for index, word in enumerate(token.split()):
        low = word.lower()
        if index and low in small:
            parts.append(low)
        else:
            parts.append(low[:1].upper() + low[1:] if low else low)
    return " ".join(parts)


def parse_place(value: str, counter: Counter) -> str:
    token = text(value)
    if not token:
        return "Unknown"
    alias = PLACE_ALIASES.get(token.lower())
    if alias:
        if alias.lower() != token.lower():
            counter[f"{token} → {alias}"] += 1
        return alias
    return pretty(token)


def parse_program(project: str, fallback_program: str, fallback_system: str) -> tuple[str, str]:
    token = text(project).lower()
    if token in PROGRAMS:
        return PROGRAMS[token]
    if token:
        return pretty(project), fallback_system
    return fallback_program, fallback_system


def parse_role(raw: str) -> tuple[str, str]:
    key = re.sub(r"[^a-z0-9& ]", " ", text(raw).lower())
    key = " ".join(key.split())
    for needle, role, family in ROLE_RULES:
        if needle in key:
            return role, family
    if not key:
        return "Unknown", "Unknown"
    return pretty(raw), "Other"


def qual_tokens(raw: str) -> list[str]:
    tokens = re.findall(r"[a-z]+", text(raw).lower())
    merged: list[str] = []
    index = 0
    shorts = {"a", "sc", "ed", "com", "s", "phil", "cs"}
    while index < len(tokens):
        current = tokens[index]
        if len(current) == 1 and index + 1 < len(tokens) and tokens[index + 1] in shorts:
            merged.append(current + tokens[index + 1])
            index += 2
            continue
        if len(current) == 1:
            buffer = current
            index += 1
            while index < len(tokens) and len(tokens[index]) == 1:
                buffer += tokens[index]
                index += 1
            merged.append(buffer)
            continue
        merged.append(current)
        index += 1
    return merged


def qualification_level(raw: str) -> str:
    cleaned = text(raw)
    tokens = set(qual_tokens(cleaned))
    if not tokens or tokens <= NONE or tokens <= {"n", "a"}:
        return "Not recorded"
    if tokens & ILLITERATE:
        return "None"
    if "phd" in tokens or "doctorate" in tokens:
        return "Doctorate"
    if tokens & MASTERS or any(token.startswith("master") for token in tokens):
        return "Masters"
    if tokens & BACHELORS or any(token.startswith("bachelor") for token in tokens):
        return "Bachelors"
    if tokens & INTER or any("inter" in token for token in tokens):
        return "Intermediate"
    if any("matric" in token for token in tokens):
        return "Matric"
    if tokens & {"middle", "midle"}:
        return "Middle"
    if tokens & {"primary", "primery"}:
        return "Primary"
    if tokens & ILLITERATE:
        return "None"
    if tokens & NONE:
        return "Not recorded"
    return "Other"


def parse_gender(value) -> str:
    token = text(value).lower()
    if token in {"f", "female", "girl", "girls"}:
        return "Female"
    if token in {"m", "male", "boy", "boys"}:
        return "Male"
    return "Unknown"


def parse_religion(value, present: bool) -> str:
    if not present:
        return "Unknown"
    token = text(value)
    if not token:
        return "Unknown"
    mapped = RELIGION.get(token.lower())
    return mapped or pretty(token)


def parse_donor(value) -> str:
    token = text(value)
    if not token:
        return "Unknown"
    mapped = DONORS.get(token.lower())
    return mapped or pretty(token)


def parse_contact(value) -> str:
    number = digits(value)
    if len(number) == 10 and number.startswith("3"):
        number = "0" + number
    if len(number) < 7:
        return ""
    return number


def parse_cnic(value) -> str:
    number = digits(value)
    if len(number) < 5:
        return ""
    if len(number) == 13:
        return f"{number[:5]}-{number[5:12]}-{number[12]}"
    return number


def column_index(headers: list[str], test) -> int | None:
    for index, header in enumerate(headers):
        if test(header):
            return index
    return None


def locate(headers: list[str]) -> dict[str, int | None]:
    return {
        "kinCnic": column_index(headers, lambda h: "cnic" in h and ("father" in h or "husband" in h)),
        "cnic": column_index(headers, lambda h: h == "cnic"),
        "name": column_index(headers, lambda h: "staff name" in h or h in {"name", "employee name"}),
        "father": column_index(headers, lambda h: ("father" in h or "husband" in h) and "cnic" not in h),
        "gender": column_index(headers, lambda h: h == "gender" or h.startswith("gender")),
        "religion": column_index(headers, lambda h: "religion" in h),
        "role": column_index(headers, lambda h: "professional" in h or h == "designation" or h == "role"),
        "qualification": column_index(headers, lambda h: "qualif" in h),
        "contact": column_index(headers, lambda h: "contact" in h or h in {"phone", "mobile"}),
        "emis": column_index(headers, lambda h: h == "emis"),
        "school": column_index(headers, lambda h: "school" in h),
        "district": column_index(headers, lambda h: h == "district"),
        "tehsil": column_index(headers, lambda h: h == "tehsil"),
        "province": column_index(headers, lambda h: h == "province"),
        "project": column_index(headers, lambda h: h == "project"),
        "donor": column_index(headers, lambda h: h == "donor"),
        "year": column_index(headers, lambda h: h == "year"),
    }


def cell(row: tuple, index: int | None):
    if index is None or index >= len(row):
        return None
    return row[index]


def main() -> None:
    if not WORKBOOK.exists():
        raise SystemExit(f"Workbook not found: {WORKBOOK}")

    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    available = {name.strip(): name for name in book.sheetnames}
    people = []
    sheet_meta = []
    places: Counter = Counter()
    roles: Counter = Counter()
    levels: Counter = Counter()
    others: Counter = Counter()

    for spec in SHEETS:
        sheet_name = available.get(spec["name"].strip())
        if not sheet_name:
            raise SystemExit(f"Missing sheet {spec['name']!r}. Found {book.sheetnames}")
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows)
        headers = [norm_header(item) for item in header_row]
        columns = locate(headers)
        if columns["name"] is None:
            raise SystemExit(f"{sheet_name}: no staff name column in {headers}")
        blank = 0
        included = 0
        seen = 0
        for raw in rows:
            seen += 1
            name = text(cell(raw, columns["name"]))
            school = text(cell(raw, columns["school"]))
            if not name and not school:
                blank += 1
                continue
            if not name:
                blank += 1
                continue
            program, system = parse_program(text(cell(raw, columns["project"])), spec["program"], spec["system"])
            role, family = parse_role(text(cell(raw, columns["role"])))
            qualification_raw = text(cell(raw, columns["qualification"]))
            level = qualification_level(qualification_raw)
            roles[f"{text(cell(raw, columns['role']))} → {role} ({family})"] += 1
            levels[level] += 1
            if level == "Other":
                others[qualification_raw or "(blank)"] += 1
            person = {
                "name": pretty(name),
                "father": pretty(text(cell(raw, columns["father"]))) if text(cell(raw, columns["father"])) else "",
                "gender": parse_gender(cell(raw, columns["gender"])),
                "religion": parse_religion(cell(raw, columns["religion"]), columns["religion"] is not None),
                "role": role,
                "family": family,
                "qualification": level,
                "qualificationRaw": qualification_raw,
                "cnic": parse_cnic(cell(raw, columns["cnic"])),
                "kinCnic": parse_cnic(cell(raw, columns["kinCnic"])),
                "contact": parse_contact(cell(raw, columns["contact"])),
                "emis": text(cell(raw, columns["emis"])),
                "school": pretty(school) if school else "Unknown",
                "district": parse_place(text(cell(raw, columns["district"])), places),
                "tehsil": parse_place(text(cell(raw, columns["tehsil"])), places),
                "province": parse_place(text(cell(raw, columns["province"])), places),
                "program": program,
                "system": system,
                "donor": parse_donor(cell(raw, columns["donor"])),
                "year": text(cell(raw, columns["year"])),
                "source": spec["source"],
            }
            if person["father"] == "Unknown":
                person["father"] = ""
            people.append(person)
            included += 1
        sheet_meta.append(
            {
                "sheet": sheet_name.strip(),
                "rows": seen,
                "blankRows": blank,
                "included": included,
                "note": spec["note"],
            }
        )

    cnic_counts = Counter(person["cnic"] for person in people if person["cnic"])
    duplicate_ids = sum(count - 1 for count in cnic_counts.values() if count > 1)
    payload = {
        "meta": {
            "staff": len(people),
            "generatedOn": date.today().isoformat(),
            "generatedFrom": WORKBOOK.name,
            "sheets": sheet_meta,
            "duplicateCnicExtra": duplicate_ids,
            "placeAliases": [{"change": change, "count": count} for change, count in places.most_common()],
            "notes": [
                "Six staff sheets are separate school lists, so every named person is included.",
                "Roles were folded into a title and a family: Teaching, Leadership, Administration, or Support.",
                "Qualifications were mapped to a level. The wording from the sheet is kept on each person.",
                "Contact numbers are digits only. A 10-digit mobile starting with 3 is shown with a leading 0.",
                "Sindh staff have no religion column, so religion stays Unknown there.",
                "The Punjab adopted sheet has no year column.",
            ],
        },
        "staff": people,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"wrote {len(people)} staff to {OUT} ({OUT.stat().st_size} bytes)")
    print("sheets", sheet_meta)
    print("levels", levels.most_common())
    print("other qualifications", others.most_common(30))
    print("unmapped roles", [item for item in roles.most_common() if "Other" in item[0] or "Unknown" in item[0]][:20])
    print("duplicate cnic extras", duplicate_ids)


if __name__ == "__main__":
    main()
